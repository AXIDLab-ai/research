"""Private Excel ingestion; public synthetic data; empirical transition/emission banks."""
from pathlib import Path
import hashlib
import json
import numpy as np
import pandas as pd

SECTORS = ["C", "J", "M", "G", "Other"]


def bootstrap_bundle(bundle, index, seed=20260920):
    """Operator-cluster bootstrap preserves firm histories and assigns fresh copy IDs."""
    operators = sorted(bundle.panel.op.unique())
    rng = np.random.default_rng(np.random.SeedSequence([seed, 9101, index]))
    samples, next_id = [], 0
    for new_op, source_op in enumerate(rng.choice(operators, len(operators), replace=True)):
        frame = bundle.panel[bundle.panel.op.eq(source_op)].copy()
        unique = sorted(frame.id.unique())
        mapping = {old: next_id + j for j, old in enumerate(unique)}
        frame["id"] = frame.id.map(mapping)
        frame["op"] = new_op
        next_id += len(unique)
        samples.append(frame)
    metadata = {**bundle.metadata, "bootstrap": index, "bootstrap_seed": seed, "parent_hash": bundle.hash}
    return Bundle(pd.concat(samples, ignore_index=True), metadata, bundle.threshold, bundle.metadata.get("strict_zero", False))


def states(frame, threshold=1000.0, strict_zero=False):
    known = frame.rev.notna() & frame.profit.notna()
    if strict_zero:
        known &= ~frame.triplezero
    z = pd.Series(np.nan, index=frame.index)
    z.loc[known] = (frame.loc[known, "rev"].ge(threshold).astype(int)
                    + 2 * frame.loc[known, "profit"].gt(0).astype(int))
    z.loc[frame.closed] = 4
    return z


def prepare_excel(source, destination, strict_zero=False):
    import openpyxl
    if hasattr(source, "getvalue"):
        source_hash = hashlib.sha256(source.getvalue()).hexdigest()
        source.seek(0)
    else:
        source = Path(source)
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    destination = Path(destination) if destination is not None else None
    if destination:
        destination.mkdir(parents=True, exist_ok=True)
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    rows = book["창업기업RAW"].iter_rows(min_row=3, max_col=openpyxl.utils.column_index_from_string("JP"), values_only=True)
    headers = next(rows)
    raw = pd.DataFrame([r for r in rows if r[0] is not None],
                       columns=[openpyxl.utils.get_column_letter(i + 1) for i in range(len(headers))])
    book.close()
    num = lambda c: pd.to_numeric(raw[c], errors="coerce")
    opmap = {v: i for i, v in enumerate(sorted(raw.C.fillna("UNKNOWN").astype(str).unique()))}
    closure = pd.to_datetime(raw.L.astype(str).str.replace(r"\.0$", "", regex=True), format="%Y%m%d", errors="coerce").dt.year.where(raw.K.eq("폐업자"))
    base = pd.DataFrame(dict(id=np.arange(len(raw)), co=num("P"), op=raw.C.fillna("UNKNOWN").astype(str).map(opmap),
                             sector=raw["T"].where(raw["T"].isin(SECTORS), "Other"), closure_year=closure,
                             closure_recorded=raw.K.eq("폐업자"), suspended=raw.K.eq("휴업자"),
                             ma=raw.J.fillna("").astype(str).str.contains("M&A"),
                             listed=raw.AD.isin(["코스닥시장", "코넥스"]), flag=num("M"), respondent=num("O")))
    parts = []
    for year in range(2013, 2025):
        f = base.copy()
        f["year"], f["age"] = year, year - f.co
        for name, start in (("rev", "GO"), ("profit", "IS"), ("assets", "JE"), ("employment", "FB")):
            f[name] = num(openpyxl.utils.get_column_letter(openpyxl.utils.column_index_from_string(start) + year - 2013))
        f["triplezero"] = f[["rev", "profit", "assets"]].eq(0).all(axis=1)
        f["closed"] = f.closure_year.le(year)
        parts.append(f)
    panel = pd.concat(parts, ignore_index=True)
    audit = {"source_sha256": source_hash, "firms": len(base),
             "missing_selection_year": int(base.co.isna().sum()),
             "closure_without_date": int((base.closure_recorded & base.closure_year.isna()).sum()),
             "preselection_triple_zero": int((panel.age.lt(0) & panel.triplezero).sum()),
             "postselection_triple_zero": int((panel.age.ge(0) & panel.triplezero).sum()),
             "closure_year_nonzero_revenue": int((panel.year.eq(panel.closure_year) & panel.rev.gt(0)).sum()),
             "privacy": "PRIVATE: pseudonyms are not anonymisation; never publish this bundle",
             "units": "KRW million, nominal; age is calendar-year offset", "synthetic": False,
             "strict_zero": strict_zero}
    panel = panel[panel.age.ge(0)].copy()
    panel["state"] = states(panel, strict_zero=strict_zero)
    if destination:
        panel.to_csv(destination / "panel.csv", index=False)
        (destination / "metadata.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
        pd.DataFrame({"column": raw.columns, "header": list(headers)}).to_csv(destination / "source_columns.csv", index=False)
        return audit
    return panel, audit


def synthetic_panel(seed=20260920, n=1200):
    """Synthetic fixture only: no copying, resampling or inference from private firms."""
    rng = np.random.default_rng(seed)
    records = []
    for i in range(n):
        sector = SECTORS[i % 5]
        co = 2013 + i % 10
        op = i % 24
        state = int(rng.choice(4, p=[.7, .15, .1, .05]))
        exit_year = np.nan
        for age in range(2025 - co):
            if age and state != 4:
                base = np.array([.5, .19, .15, .13, .03])
                base[state] += .7
                base[3] += .025 * min(age, 5)
                state = int(rng.choice(5, p=base / base.sum()))
            if state == 4 and np.isnan(exit_year):
                exit_year = co + age
            z = rng.normal()
            r = (1000 + np.exp(7 + z)) if state % 2 else (0.0 if rng.random() < .15 else 999 * rng.beta(1.2, 2))
            p = (1 if state in (2, 3) else -1) * np.exp(4.5 + .7 * z + .4 * rng.normal())
            if state == 4 or rng.random() < .08:
                r = p = np.nan
            records.append(dict(id=i, co=co, op=op, sector=sector, year=co + age, age=age,
                                rev=r, profit=p, assets=r * 2, closed=state == 4, triplezero=False,
                                closure_year=exit_year, closure_recorded=state == 4, state=state))
    panel = pd.DataFrame(records)
    # Survey history may be missing for firms recorded closed at the survey date.
    closed_ids = panel.loc[panel.closed, "id"].unique()
    remove = panel.id.isin(closed_ids) & (rng.random(len(panel)) < .8)
    panel.loc[remove, ["rev", "profit", "assets"]] = np.nan
    panel["state"] = states(panel)
    return panel


class Bundle:
    def __init__(self, panel, metadata, threshold=1000.0, strict_zero=False):
        self.panel = panel.copy().sort_values(["id", "age"])
        self.metadata = {**metadata, "strict_zero": strict_zero}
        self.threshold = threshold
        self.panel["state"] = states(self.panel, threshold, strict_zero)
        # Only the declared calibration split estimates kernels and emissions.
        train = self.panel[self.panel.co.le(2019) & self.panel.age.between(0, 2)].copy()
        self.train = train
        self.sectors = SECTORS
        counts = train[train.age.eq(0)].sector.value_counts().reindex(SECTORS, fill_value=0).to_numpy(float)
        self.weights = (counts + .5) / (counts.sum() + 2.5)
        self.operator_sizes = train[train.age.eq(0)].groupby("op").size().to_numpy(float)
        if not len(self.operator_sizes):
            raise ValueError("No entry cohorts in calibration split")
        self.initial = np.zeros((5, 4))
        self.transition = np.zeros((5, 5, 5))
        self.banks = {}
        previous = train.groupby("id")[["age", "state"]].shift()
        train["previous"] = previous.state
        pairs = train[train.age.eq(previous.age + 1) & train.state.notna() & previous.state.notna()].copy()
        # Exclude ambiguous positive financial values in recorded closure year.
        pairs = pairs[~(pairs.state.eq(4) & (pairs.rev.gt(0) | pairs.profit.ne(0) & pairs.profit.notna()))]
        pooled = np.full((5, 5), .5)
        for row in pairs.itertuples():
            pooled[int(row.previous), int(row.state)] += 1
        pooled /= pooled.sum(axis=1, keepdims=True)
        for s, sector in enumerate(SECTORS):
            g = train[train.sector.eq(sector)]
            ini = g[g.age.eq(0) & g.state.lt(4)].state.value_counts().reindex(range(4), fill_value=0).to_numpy(float) + .5
            self.initial[s] = ini / ini.sum()
            local = 10 * pooled.copy()
            for row in pairs[pairs.sector.eq(sector)].itertuples():
                local[int(row.previous), int(row.state)] += 1
            local /= local.sum(axis=1, keepdims=True)
            local[4] = [0, 0, 0, 0, 1]
            self.transition[s] = local
            for k in range(4):
                bank = g[g.state.eq(k)][["rev", "profit"]].dropna().to_numpy(float)
                if len(bank) < 5:
                    bank = train[train.state.eq(k)][["rev", "profit"]].dropna().to_numpy(float)
                if not len(bank):
                    raise ValueError(f"No empirical emission for state {k}; supply more data or use synthetic mode")
                self.banks[s, k] = bank[np.argsort(bank[:, 1], kind="stable")]
        latest = self.panel.groupby("id").tail(1).set_index("id").closed
        self.survey_missing = {}
        for closed in (False, True):
            g = train[train.id.map(latest).eq(closed)]
            self.survey_missing[closed] = float((g.rev.isna() | g.profit.isna()).mean()) if len(g) else 0.1
        serialized = self.panel.to_csv(index=False).encode()
        self.hash = hashlib.sha256(serialized).hexdigest()

    @classmethod
    def load(cls, path=None, threshold=1000.0, strict_zero=False):
        if path is None or str(path) == "synthetic":
            return cls(synthetic_panel(), {"synthetic": True, "units": "arbitrary monetary units", "seed": 20260920}, threshold, strict_zero)
        path = Path(path)
        return cls(pd.read_csv(path / "panel.csv"), json.loads((path / "metadata.json").read_text(encoding="utf-8")), threshold, strict_zero)

    def emit(self, sectors, states_, u, v, family="joint"):
        r, p = np.zeros(len(u)), np.zeros(len(u))
        for s in range(5):
            for k in range(4):
                ix = np.where((sectors == s) & (states_ == k))[0]
                if not len(ix):
                    continue
                bank = self.banks[s, k]
                a = np.minimum((u[ix] * len(bank)).astype(int), len(bank) - 1)
                b = np.minimum((v[ix] * len(bank)).astype(int), len(bank) - 1)
                r[ix], p[ix] = bank[a, 0], bank[a if family == "joint" else b, 1]
                if family == "lognormal":
                    from statistics import NormalDist
                    # Empirical sign/zero masses retained; positive magnitudes approximated.
                    for out, col, positive in ((r, 0, k % 2 == 1), (p, 1, k >= 2)):
                        vals = np.abs(bank[:, col]); nz = vals[vals > 0]
                        if not len(nz):
                            out[ix] = 0
                            continue
                        logv = np.log(nz)
                        q = np.array([NormalDist().inv_cdf(float(x)) for x in np.clip(v[ix], 1e-9, 1 - 1e-9)])
                        draw = np.exp(np.clip(logv.mean() + logv.std() * q, -30, 30))
                        draw[u[ix] < np.mean(vals == 0)] = 0
                        if col == 0:
                            draw = np.maximum(draw, self.threshold) if positive else np.minimum(draw, self.threshold * (1 - 1e-9))
                        else:
                            draw = np.maximum(draw, 1e-9) if positive else -draw
                        out[ix] = draw
        return r, p
